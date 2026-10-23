"""Aplica a la copia de la memoria las correcciones contrastadas con el código.

La fuente Word original queda intacta. Las aserciones detienen la generación si
se intenta aplicar el guion a una versión distinta de la revisada.
"""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/memoria/TFM_version_tutor_2026-10.docx"
OUTPUT = ROOT / "docs/memoria/TFM_revision_coherencia_2026-10.docx"


def main() -> None:
    document = Document(SOURCE)

    def paragraph(index: int, starts: str, replacement: str) -> None:
        item = document.paragraphs[index]
        if not item.text.startswith(starts):
            raise ValueError(f"La memoria cambió en el párrafo {index}: {item.text[:80]}")
        if not item.runs:
            item.add_run()
        item.runs[0].text = replacement
        for run in item.runs[1:]:
            run.text = ""

    def replace(index: int, original: str, replacement: str) -> None:
        item = document.paragraphs[index]
        hits = [run for run in item.runs if original in run.text]
        if len(hits) != 1:
            raise ValueError(f"Se esperaban una coincidencia en {index}: {original}")
        hits[0].text = hits[0].text.replace(original, replacement, 1)

    paragraph(
        18,
        "Este Trabajo Fin de Máster presenta",
        "Este Trabajo Fin de Máster presenta un sistema de apoyo a la decisión "
        "para estimar necesidades de riego diarias y semanales. La visión del "
        "producto abarca distintos cultivos y zonas, pero el prototipo y su "
        "evaluación se limitan al pimiento y a la estación AL01 La Mojonera "
        "(Almería), durante la campaña de 2025. La ETL integra observaciones y "
        "necesidades de SiAR; una predicción AEMET de 2026 prueba la integración, "
        "pero no alimenta el motor evaluado. De 150 fechas observadas, 148 "
        "tienen una referencia diaria válida. Seis enfoques se compararon con "
        "validación temporal: ningún algoritmo de aprendizaje superó de forma "
        "estable a la persistencia. Una vía agronómica que combina ET0 reciente "
        "y Kc redujo el MAE a un día un 7 % en desarrollo y un 25 % en el tramo "
        "de prueba reutilizado. Son resultados frente al cálculo de SiAR, no "
        "demuestran ahorro de agua ni validez para una parcela sin contrastación "
        "agronómica y evaluación en campo.",
    )
    paragraph(
        170,
        "A partir de este contexto",
        "El anteproyecto del TFM formuló la visión de un sistema inteligente "
        "de apoyo a la decisión para la gestión hídrica agrícola, aplicable "
        "a distintos cultivos y localizaciones. Esta memoria distingue esa "
        "arquitectura objetivo del piloto construido: un flujo reproducible "
        "de observaciones SiAR, comparación predictiva y recomendación para "
        "pimiento en AL01 La Mojonera durante 2025. Los indicadores "
        "hidrológicos y las predicciones AEMET se estudiaron o extrajeron, "
        "pero no sustentan la evaluación del motor. La interfaz facilita "
        "la interpretación de resultados orientativos; no sustituye la "
        "decisión profesional ni determina por sí sola cuándo regar.",
    )
    for item in document.paragraphs:
        for run in item.runs:
            if "5. Sistema de optimización y resultados" in run.text:
                run.text = run.text.replace(
                    "5. Sistema de optimización y resultados",
                    "5. Sistema de recomendación y resultados",
                )
    toc_chapter_5 = document.paragraphs[73]
    if "Sistema de optimización y resultados" not in toc_chapter_5.text:
        raise ValueError("No se encontró el capítulo 5 en el índice estático")
    toc_chapter_5.text = toc_chapter_5.text.replace(
        "Sistema de optimización y resultados",
        "Sistema de recomendación y resultados",
    )

    replace(
        174,
        "que integre datos meteorológicos, ambientales, agronómicos e hidrológicos, estime las necesidades de riego de los cultivos y las transforme en recomendaciones operativas, trazables y comprensibles para las localizaciones incluidas en el alcance del estudio.",
        "que convierta datos agroclimáticos y parámetros de parcela en recomendaciones diarias y semanales comprensibles y trazables. La visión del producto comprende varios cultivos y zonas; el objetivo evaluable de este TFM se restringe a un piloto de una estación, un cultivo y una campaña. La integración hidrológica automática y la decisión sobre el momento de regar requieren datos adicionales.",
    )
    paragraph(
        184,
        "El trabajo es una prueba de concepto",
        "El trabajo es una prueba de concepto de un sistema de apoyo a la "
        "decisión que conecta adquisición de datos, análisis, comparación de "
        "modelos y recomendación. La evaluación empírica se limita al pimiento "
        "en la estación AL01 La Mojonera, entre mayo y septiembre de 2025. El "
        "diseño permite preparar otros archivos de estación y cultivo, pero la "
        "interfaz presentada está fijada al caso piloto y no permite elegirlos. "
        "Los resultados no se extrapolan a otros casos.",
    )
    paragraph(
        202,
        "Estos fundamentos proporcionan",
        "Estos fundamentos proporcionan la referencia agronómica del sistema. "
        "El piloto utiliza la evapotranspiración de referencia, el coeficiente "
        "del cultivo y la precipitación efectiva para aproximar la necesidad "
        "neta diaria; no dispone de humedad del suelo ni resuelve su balance "
        "hídrico. Por ello, la aplicación calcula una dosis orientativa y no "
        "determina de forma autónoma el momento óptimo de riego. Los modelos "
        "de aprendizaje automático se evalúan como alternativas predictivas, "
        "sin sustituir la validación agronómica (Allen et al., 1998).",
    )
    replace(
        250,
        "Los archivos raw y las tablas completas generadas se mantienen fuera del control de versiones; se versionan código, esquemas, metadatos, muestras ligeras, métricas y figuras reproducibles.",
        "Los archivos raw y las tablas completas generadas se excluyen por defecto del control de versiones. Como excepción documentada, se conserva una instantánea pública y acotada de las entradas del piloto para reproducirlo sin tokens. Se versionan además código, esquemas, metadatos, métricas y figuras reproducibles.",
    )
    replace(
        408,
        "Las dos filas WARN presentan ausencia de ET0 y de las magnitudes derivadas; se mantienen para trazabilidad, pero deberán excluirse o tratarse dentro de cada partición de entrenamiento.",
        "Las dos filas WARN (21 y 22 de julio) carecen de ET0 y de las magnitudes derivadas. La del 21 también carece de temperatura y humedad medias, por lo que estas variables tienen 149 valores válidos en la Tabla 4.11. Los registros se conservan para trazabilidad y se excluyen del entrenamiento cuando falta la variable objetivo.",
    )
    paragraph(
        358,
        "No se combinan todavía",
        "No se combinaron las observaciones SiAR de 2025 con las predicciones "
        "AEMET emitidas en 2026. Una observación describe lo que ocurrió y "
        "una predicción lo que se esperaba en el momento de emisión. Para "
        "una evaluación conjunta futura harán falta pronósticos históricos "
        "y una unión por fecha válida, emisión y horizonte.",
    )
    paragraph(
        366,
        "Ejemplo de necesidades hídricas",
        "Ejemplo de necesidades hídricas del pimiento: cuatro fechas mensuales "
        "seleccionadas; no son filas consecutivas.",
    )
    paragraph(
        369,
        "Ejemplo de la predicción diaria",
        "Ejemplo de predicción diaria AEMET: se muestran D+0, D+1, D+2 y "
        "D+4; el horizonte D+3 se omite solo en esta muestra.",
    )
    replace(437, "Este control será la base para seleccionar", "Este control sirvió para seleccionar")
    replace(
        440,
        "En este piloto se propone entrenar con el tramo inicial y evaluar con fechas posteriores;",
        "En este piloto se entrenó con el tramo inicial y se evaluó con fechas posteriores;",
    )
    paragraph(
        185,
        "Las fuentes principales previstas",
        "La fuente de la evaluación es SiAR: sus observaciones agroclimáticas "
        "y el CSV de necesidades del pimiento. AEMET OpenData se extrajo y "
        "normalizó como prueba de integración, pero sus predicciones de 2026 "
        "no coinciden con la campaña de 2025 y no se utilizaron para entrenar "
        "ni validar el motor. Los indicadores agregados del Boletín Hidrológico "
        "del MITECO se estudiaron y se descartaron como asignación de agua a "
        "una parcela; la dotación la introduce el usuario.",
    )
    paragraph(
        186,
        "El horizonte operativo será",
        "El prototipo emite una recomendación diaria o una planificación de "
        "hasta siete días. Calcula necesidad neta, dosis bruta según eficiencia "
        "y volumen cuando se indica la superficie. Sin balance hídrico del "
        "suelo y registro de riegos, no determina el momento óptimo de regar; "
        "una lluvia prevista solo produce un aviso. Tampoco acciona bombas, "
        "válvulas ni programadores.",
    )
    paragraph(
        188,
        "La interacción prevista permitirá",
        "La interfaz implementada permite elegir fecha, horizonte, sistema y "
        "eficiencia de riego, superficie y dotación para el caso piloto. La "
        "selección de cultivo y localización pertenece a la visión del producto, "
        "no a esta versión. La decisión final permanece en manos del usuario, "
        "que debe contrastarla con el estado real de la parcela y el criterio "
        "agronómico de un profesional.",
    )
    replace(
        227,
        "El producto se ha diseñado para que el usuario seleccione ubicación, cultivo y periodo.",
        "La visión del producto contempla que el usuario seleccione ubicación, cultivo y periodo; esa interfaz no se ha implementado en el piloto.",
    )
    replace(
        227,
        "su tratamiento como aire libre es una decisión de alcance del piloto, coherente con el uso de meteorología y precipitación exterior.",
        "su tratamiento como aire libre es una hipótesis de alcance del piloto, coherente con el uso de meteorología y precipitación exterior, pero pendiente de acreditación para el CSV concreto empleado.",
    )
    paragraph(
        286,
        "Las credenciales se almacenan",
        "Las credenciales se conservan en .env, excluido de Git. Las capas raw, "
        "interim y processed también se excluyen por defecto. Una excepción "
        "acotada de datos públicos del piloto permite reconstruir la ETL sin "
        "tokens; se documenta en docs/REPRODUCCION_PILOTO.md. Se versionan "
        "asimismo código, métricas y figuras.",
    )
    replace(274, "cincuenta y cinco pruebas", "sesenta y una pruebas")
    replace(
        290,
        "Las predicciones se obtienen para el municipio AEMET 04013 Almería.",
        "Se obtuvo una predicción de prueba para el municipio AEMET 04013 Almería en 2026; no alimenta la evaluación retrospectiva de 2025.",
    )
    replace(
        290,
        "El piloto utiliza como referencia el pimiento al aire libre en Almería.",
        "El piloto plantea como referencia el pimiento al aire libre en Almería, con la correspondencia del CSV de SiAR a ese sistema aún pendiente de acreditar.",
    )
    replace(
        290,
        "El diseño conserva identificadores de fuente, estación, municipio, cultivo y tiempo para que la selección del usuario determine qué datos se consultan en cada ejecución.",
        "El diseño conserva identificadores de fuente, estación, municipio, cultivo y tiempo para preparar otros casos; la interfaz actual no ofrece esa selección.",
    )
    replace(
        289,
        "El sistema de apoyo a la decisión necesita combinar observaciones agroclimáticas históricas con predicciones meteorológicas.",
        "La visión del sistema combina observaciones agroclimáticas con predicciones meteorológicas; el piloto evaluado utiliza las primeras y deja las segundas como prueba de integración.",
    )
    paragraph(
        341,
        "La partición por fuente",
        "La partición por fuente, producto, estación o municipio y fecha permite "
        "cargas incrementales. Los raw de trabajo no se publican por defecto. "
        "Como excepción se versiona una instantánea pública del piloto con "
        "atribución a SiAR y AEMET, suficiente para reconstruir la ETL, la "
        "comparación de modelos y la validación del sistema sin credenciales.",
    )
    paragraph(
        343,
        "El segundo proceso selecciona",
        "El segundo proceso genera tablas planas. Cuando existen varios "
        "históricos SiAR o CSV de cultivo exige que se indiquen las rutas, en "
        "vez de elegir el archivo más reciente. Antes de integrar comprueba "
        "la coincidencia de estación y la intersección de fechas. Las reglas "
        "implementadas son:",
    )
    paragraph(
        456,
        "Los nueve predictores utilizados",
        "La tabla analítica tiene 88 columnas, pero muchas son identificadores, "
        "controles de calidad, variables del propio objetivo o meteorología del "
        "día futuro no disponible al emitir. Los nueve predictores seguros son "
        "Kc, las dos componentes cíclicas del día del año, tres retardos de la "
        "necesidad y tres acumulados históricos. Se calculan por fecha de "
        "calendario, de modo que los huecos no se tratan como días consecutivos.",
    )
    replace(
        459,
        "bloque inmediatamente posterior.",
        "bloque inmediatamente posterior. Los tres cortes usan 32/31, 63/31 y 94/31 observaciones de entrenamiento/validación y representan predicciones a un día con nuevas observaciones disponibles, no una semana futura completa.",
    )
    paragraph(
        460,
        "Este diseño es más robusto",
        "La selección de los seis enfoques se hizo por el MAE medio de los tres "
        "cortes de desarrollo antes de consultar los 23 registros de septiembre. "
        "El capítulo 5 volvió a usar esas fechas para evaluar el método ET0 "
        "reciente por Kc; por ello no deben presentarse como una segunda prueba "
        "independiente de todo el proyecto. Se fija la semilla 42 y las "
        "predicciones se limitan a un mínimo físico de 0 mm/día.",
    )
    replace(
        471,
        "Esta última métrica es coherente con el DSS",
        "Aquí los bloques semanales parten de la primera fecha de cada corte y suman hasta siete días de calendario, incluido un bloque final parcial; no equivalen a las semanas naturales completas del apartado 5.3.5. Esta métrica es coherente con el DSS",
    )
    replace(
        511,
        "La vía preferente utiliza una ET0 prevista cuando el usuario o la fuente la proporcionan, ya sea el producto pronosticado de SiAR o una estimación calculada a partir de la predicción de AEMET.",
        "La vía preferente utiliza una ET0 prevista solo cuando se proporciona como entrada explícita. El prototipo no descarga ni calcula automáticamente ese pronóstico; los productos de SiAR y una estimación derivada de AEMET son posibilidades futuras pendientes de evaluación.",
    )
    replace(
        520,
        "Los umbrales de horizonte se apoyan en los errores medidos en el apartado 5.3 y pueden ajustarse cuando se disponga de más campañas.",
        "El cuarto día es un umbral prudencial de interfaz, no una probabilidad calibrada. Los MAE por horizonte del apartado 5.3 describen este piloto y no son intervalos de predicción individuales.",
    )
    replace(
        559,
        "Para mantener la coherencia con el capítulo 4, se utilizó como prueba el mismo tramo reservado, del 8 al 30 de septiembre, y la selección se hizo únicamente con el periodo de desarrollo.",
        "Se reutilizó el tramo del 8 al 30 de septiembre para la evaluación, aunque la ventana se seleccionó únicamente con el periodo de desarrollo. Estas fechas ya habían sido consultadas en el capítulo 4 y no constituyen una segunda prueba independiente.",
    )
    replace(
        567,
        "El sistema reduce el MAE respecto a la persistencia",
        "Los 124 días comparables de desarrollo son uno menos que las 125 etiquetas del capítulo 4: la primera emisión se realiza el 2 de mayo y necesita historia previa. El sistema reduce el MAE respecto a la persistencia",
    )
    replace(
        573,
        "habitual en predicción de series temporales (Hyndman & Athanasopoulos, 2021).",
        "habitual en predicción de series temporales (Hyndman & Athanasopoulos, 2021). La simulación empieza el 16 de mayo para disponer de historia suficiente en todas las ventanas. Por eso su día 1 de desarrollo y la Tabla 5.4 usan muestras diferentes; en prueba coinciden las 23 emisiones.",
    )
    paragraph(
        579,
        "Estos resultados justifican",
        "El horizonte de siete días se ofrece como planificación orientativa, "
        "aunque el séptimo día de prueba pierde frente a la persistencia. "
        "Asignar confianza baja desde el cuarto día es una precaución de la "
        "interfaz, no un umbral calibrado estadísticamente. La regla R9 muestra "
        "el MAE histórico por día del horizonte; no representa un intervalo "
        "individual ni garantiza ese error en otra campaña. Para anticipar "
        "cambios meteorológicos se necesita una ET0 prevista validada.",
    )
    paragraph(
        587,
        "La eficiencia es el parámetro",
        "La eficiencia modifica la dosis bruta. La Tabla 5.6 utiliza los "
        "531,05 mm netos de los 148 días con referencia válida. No imputa las "
        "tres fechas ausentes ni los dos registros sin ET0, por lo que no "
        "representa el total íntegro de los 153 días de calendario.",
    )
    for index in (142, 588):
        replace(
            index,
            "Dosis bruta de la campaña según el sistema de riego",
            "Dosis bruta de los 148 días con referencia válida según el sistema de riego",
        )
    paragraph(
        596,
        "Las comprobaciones de la Tabla 5.7",
        "Los trece escenarios se complementan con quince pruebas automáticas "
        "del motor. Tras añadir controles de selección de archivos y "
        "compatibilidad entre estación y campaña, la suite completa contiene "
        "61 pruebas. También verifica que una planificación no cambie si se "
        "alteran datos posteriores a la fecha de emisión.",
    )
    replace(613, "cadena completa y reproducible que parte de las APIs públicas de SiAR y AEMET", "cadena reproducible que parte de observaciones y necesidades de SiAR")
    replace(613, "Esa afirmación tiene un límite preciso:", "AEMET se extrajo y normalizó, pero su emisión de 2026 no alimenta el motor evaluado con 2025. Esa afirmación tiene un límite preciso:")
    replace(620, "consultar la prueba final. Leídos conjuntamente", "consultar la prueba final para esa comparación. El método ET0 reciente por Kc se evaluó después reutilizando septiembre; no es otra validación independiente. Leídos conjuntamente")
    replace(622, "cincuenta y cinco pruebas automáticas", "sesenta y una pruebas automáticas")
    replace(609, "cincuenta y cinco pruebas automáticas", "sesenta y una pruebas automáticas")
    replace(628, "a partir del cuarto día del horizonte la recomendación depende de que el tiempo se mantenga", "en cualquier horizonte futuro la recomendación depende de que el tiempo no cambie demasiado; el cuarto día solo activa un aviso prudencial")
    replace(634, "tres impactos verificables y uno potencial", "tres resultados observables en el prototipo y un impacto potencial, todavía no medido en campo")
    replace(635, "El primero es operativo.", "El primero es una capacidad funcional comprobada.")
    replace(635, "se traduce en un déficit hídrico precisamente cuando la demanda aumenta.", "podría producir un déficit si se aplicara sin otros controles; el ejemplo retrospectivo no demuestra que ocurriera en una parcela.")
    paragraph(
        642,
        "La primera mejora consiste",
        "La primera mejora consiste en ampliar los datos con varias campañas, "
        "estaciones de Almería y cultivos. La extracción meteorológica admite "
        "estación y periodo; las necesidades se aportan mediante un CSV SiAR "
        "y metadatos. La ETL ya exige selección explícita cuando hay varios "
        "archivos y rechaza estaciones o campañas incompatibles. Falta un "
        "catálogo seleccionable en la interfaz y comprobar la aplicabilidad "
        "agronómica de cada calendario. Con más campañas podrán reservarse "
        "bloques completos para evaluar generalización temporal y espacial "
        "(Roberts et al., 2017).",
    )

    changes = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if "13 escenarios y 55 pruebas automáticas" in cell.text:
                    for item in cell.paragraphs:
                        for run in item.runs:
                            if "13 escenarios y 55 pruebas automáticas" in run.text:
                                run.text = run.text.replace(
                                    "13 escenarios y 55 pruebas automáticas",
                                    "13 escenarios y 61 pruebas automáticas",
                                )
                                changes += 1
    if changes != 1:
        raise ValueError(f"No se encontró la referencia a 55 pruebas en Tabla 6.1: {changes}")

    source_table = document.tables[4]
    if source_table.cell(0, 0).text != "Organismo y producto":
        raise ValueError("No se encontró la tabla 4.1 de fuentes")
    source_table.cell(0, 1).text = "Producto consultado"
    source_header_run = source_table.cell(0, 1).paragraphs[0].runs[0]
    source_header_run.bold = True
    source_header_run.font.color.rgb = RGBColor(255, 255, 255)
    source_labels = ["MAPA · SiAR", "MAPA · calculadora SiAR", "AEMET OpenData"]
    access_labels = ["Web API SiAR v1", "Web de necesidades hídricas", "AEMET OpenData API"]
    for row_index, (source_label, access_label) in enumerate(
        zip(source_labels, access_labels), start=1
    ):
        source_table.cell(row_index, 0).text = source_label
        source_table.cell(row_index, 1).text = access_label
    for column, width in zip(source_table.columns, [1.35, 2.45, 1.35, 1.53]):
        column.width = Inches(width)
    for row in source_table.rows:
        for cell, width in zip(row.cells, [1.35, 2.45, 1.35, 1.53]):
            cell.width = Inches(width)

    source_inventory = document.tables[6]
    if source_inventory.cell(0, 2).text != "Registros":
        raise ValueError("No se encontró la cabecera de la tabla 4.3")
    source_inventory.cell(0, 2).text = "Volumen o cobertura"
    test_metrics_table = document.tables[17]
    if test_metrics_table.cell(1, 4).text != "-0,000":
        raise ValueError("No se encontró el cero negativo de la tabla 4.14")
    test_metrics_table.cell(1, 4).text = "0,000"

    weekly_example_table = document.tables[19]
    if weekly_example_table.cell(0, 0).text != "Fecha" or len(weekly_example_table.rows) != 9:
        raise ValueError("No se encontró la tabla 5.2 de planificación")
    for row in weekly_example_table.rows[:-1]:
        for cell in row.cells:
            for item in cell.paragraphs:
                item.paragraph_format.keep_with_next = True

    architecture_paragraph = document.paragraphs[254]
    if len(architecture_paragraph._element.xpath(".//w:drawing")) != 1:
        raise ValueError("No se encontró la imagen de la figura 3.1")
    architecture_width = document.inline_shapes[1].width
    architecture_paragraph.clear()
    architecture_paragraph.add_run().add_picture(
        str(ROOT / "docs/images/arquitectura_sistema_3_2.png"),
        width=architecture_width,
    )
    architecture_caption = document.paragraphs[255]
    if not architecture_caption.text.startswith("Figura 3.1."):
        raise ValueError("No se encontró el pie de la figura 3.1")
    architecture_caption.paragraph_format.keep_with_next = False
    figure_paragraph = document.paragraphs[291]
    if len(figure_paragraph._element.xpath(".//w:drawing")) != 1:
        raise ValueError("No se encontró la imagen de la figura 4.1")
    figure_width = document.inline_shapes[1].width
    figure_paragraph.clear()
    figure_paragraph.add_run().add_picture(
        str(ROOT / "docs/images/procesamiento_datos_4_1.png"), width=figure_width
    )
    paragraph(
        292,
        "Figura 4.1.",
        "Figura 4.1. ETL del caso piloto desde su configuración hasta la tabla "
        "analítica. Los gráficos inferiores utilizan datos reales de la "
        "extracción del 24 de septiembre de 2026.",
    )
    # La leyenda de la figura 4.1 cabe en la página de la figura. Su estilo
    # heredaba «mantener con el siguiente», que la arrastraba junto al título
    # de 4.1.1 y dejaba la figura sin pie.
    figure_caption = document.paragraphs[292]
    if not figure_caption.text.startswith("Figura 4.1."):
        raise ValueError("No se encontró el pie de la figura 4.1")
    figure_caption.paragraph_format.keep_with_next = False
    model_plot_paragraph = document.paragraphs[482]
    if len(model_plot_paragraph._element.xpath(".//w:drawing")) != 1:
        raise ValueError("No se encontró la figura 4.7 de predicciones")
    model_plot_paragraph.clear()
    model_plot_paragraph.add_run().add_picture(
        str(ROOT / "docs/images/model_predictions_4_3.png"), width=Inches(5.7)
    )
    first_chapter_four_paragraph = document.paragraphs[289]
    if not first_chapter_four_paragraph.text.startswith("La visión del sistema"):
        raise ValueError("No se encontró el inicio del apartado 4.1")
    first_chapter_four_paragraph.paragraph_format.space_before = Pt(12)

    project_phases_table = document.tables[3]
    if project_phases_table.cell(0, 0).text != "Fase":
        raise ValueError("No se encontró la tabla de fases del apartado 3.3")
    milestones_xml = OxmlElement("w:p")
    project_phases_table._tbl.addnext(milestones_xml)
    milestones = Paragraph(milestones_xml, project_phases_table._parent)
    milestones.style = "Normal"
    milestones.add_run(
        "El historial de Git sitúa los hitos ejecutados: documentación de "
        "fuentes y ETL el 24/09/2026; análisis y modelado entre el 28/09 y "
        "el 07/10; motor de recomendación y validación el 12/10; "
        "conclusiones entre el 15 y el 17/10; revisión de coherencia y "
        "reproducción independiente el 22/10. Son fechas de integración "
        "comprobables, no una asignación individual de tareas. La publicación "
        "de una demo por enlace y una validación agronómica externa siguen "
        "pendientes."
    )

    notation_anchor = document.paragraphs[195]
    if not notation_anchor.text.startswith("La metodología FAO-56"):
        raise ValueError("No se encontró el ancla de notación del apartado 2.1")
    notation_xml = OxmlElement("w:p")
    notation_anchor._p.addnext(notation_xml)
    notation = Paragraph(notation_xml, notation_anchor._parent)
    notation.style = "Normal"
    notation.add_run(
        "Notación empleada: ET0 es la evapotranspiración de referencia "
        "(mm/día), Kc el coeficiente de cultivo, ETc la evapotranspiración "
        "del cultivo y Pe la precipitación efectiva. En el cálculo de referencia: "
        "ETc = Kc × ET0; necesidad neta = max(ETc − Pe, 0). La dosis bruta "
        "orientativa se obtiene dividiendo la necesidad neta por la eficiencia "
        "de aplicación ηa, expresada entre 0 y 1. Estas ecuaciones no "
        "sustituyen un balance de agua del suelo."
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
