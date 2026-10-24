"""Añade a la memoria de revisión la integración real ETc + AEMET."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/memoria/TFM_revision_reentrenamiento_2026-10-23.docx"
OUTPUT = ROOT / "docs/memoria/TFM_revision_integracion_ETC_AEMET_2026-10-24.docx"


def before(doc, anchor, text: str, style: str = "Normal"):
    paragraph = doc.add_paragraph(text, style=style)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def one_heading(doc, prefix: str, style: str):
    matches = [p for p in doc.paragraphs if p.style.name == style and p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Encabezado ambiguo: {prefix}: {len(matches)}")
    return matches[0]


def add_table(doc, anchor):
    rows = [
        ["Cultivo y estación", "Modelo", "MAE modelo", "MAE físico", "Dosis"],
        ["Pimiento AL01", "Persistencia", "0,390", "0,295", "Físico"],
        ["Coliflor AL01", "Ridge", "0,762", "0,636", "Físico"],
        ["Maíz Z06", "Persistencia", "0,718", "0,556", "Físico"],
        ["Algodón BA08", "Persistencia", "0,410", "0,434", "Modelo"],
        ["Alfalfa Z06", "Persistencia", "0,556", "0,531", "Físico"],
        ["Alcachofa BA08", "Persistencia", "0,482", "0,455", "Físico"],
    ]
    table = doc.add_table(rows=len(rows), cols=5, style="TableNormal")
    table.autofit = False
    widths = [Inches(1.48), Inches(1.02), Inches(1.02), Inches(1.02), Inches(0.94)]
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.width = widths[column_index]
            for paragraph in cell.paragraphs:
                if column_index:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    run.bold = row_index == 0
    anchor._p.addprevious(table._tbl)


def repair_index(doc):
    """Actualiza los números desplazados y añade las tres entradas nuevas."""
    for paragraph in doc.paragraphs:
        if "\t" not in paragraph.text or not (
            paragraph.style.name.startswith("toc")
            or paragraph.style.name == "Entrada de índice TFM"
        ):
            continue
        title, page = paragraph.text.rsplit("\t", 1)
        if not page.isdigit():
            continue
        if title.startswith(("5.", "6.", "Anexo A.", "Tabla 5.", "Tabla 6.", "Tabla A.", "Figura 5.")):
            texts = paragraph._p.findall(".//" + qn("w:t"))
            if not texts or texts[-1].text != page:
                raise ValueError(f"Índice no actualizable: {title}")
            texts[-1].text = str(int(page) + 1)

    entries = [
        ("4.3.10. Reentrenamiento con datos de 2026", "4.3.11. Modelo de ETc para la aplicación\t54"),
        ("5.1.6. Decisiones descartadas", "5.1.7. Integración operativa de ETc y AEMET\t60"),
        ("Tabla 4.15. Reentrenamiento por cultivo", "Tabla 4.16. Modelo de ETc y cálculo físico por cultivo y estación.\t54"),
    ]
    for old, new_text in entries:
        matches = [p for p in doc.paragraphs if p.text.startswith(old) and "\t" in p.text]
        if len(matches) != 1:
            raise ValueError(f"Entrada de índice ambigua: {old}")
        anchor = matches[0]
        paragraph = doc.add_paragraph(new_text, style=anchor.style)
        for stop in anchor.paragraph_format.tab_stops:
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                stop.position, stop.alignment, stop.leader
            )
        anchor._p.addnext(paragraph._p)


def main():
    doc = Document(SOURCE)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("El capítulo anterior estableció dos referentes."):
            paragraph.text = (
                "El capítulo anterior estableció dos referentes. La necesidad neta "
                "de SiAR permite construir un cálculo agronómico explicable en días "
                "cerrados. Los modelos originales no superaron de forma estable a "
                "persistencia. La ampliación operativa del apartado 4.3.11 compara "
                "modelos de ETc con ET0 reciente por Kc y usa la base con menor error "
                "de prueba; después aplica las reglas de lluvia y dosis. No se ha "
                "validado un modelo común ni una dosis óptima en campo."
            )
        elif "Cada consulta genera una tabla diaria con dieciocho campos" in paragraph.text:
            paragraph.text = paragraph.text.replace("dieciocho campos", "veinticuatro campos")
        elif "tabla de dieciocho columnas" in paragraph.text:
            paragraph.text = paragraph.text.replace("dieciocho columnas", "veinticuatro columnas")
    chapter5 = one_heading(doc, "5. Sistema", "Heading 1")
    before(doc, chapter5, "4.3.11. Modelo de ETc para la aplicación", "Heading 3")
    before(
        doc, chapter5,
        "Los modelos anteriores anticipaban la necesidad neta calculada por SiAR, "
        "que ya incorpora precipitación efectiva. Para combinar una previsión AEMET "
        "sin descontar lluvia dos veces se entrenaron, por separado, seis modelos "
        "de ETc: demanda del cultivo antes de lluvia. Sus predictores usan Kc, "
        "estacionalidad y retardos de ETc anteriores al día objetivo. Se preservó "
        "la selección temporal y la prueba final del experimento anterior."
    )
    before(
        doc, chapter5,
        "Tabla 4.16. Modelo de ETc y cálculo físico por cultivo y estación. "
        "MAE de prueba en mm/día; valores menores indican menos error.",
        "Índice tabla TFM",
    )
    add_table(doc, chapter5)
    before(
        doc, chapter5,
        "El modelo seleccionado por validación fue persistencia en cinco casos y "
        "Ridge en coliflor. El cálculo físico que ya utilizaba el DSS —media de "
        "ET0 reciente por Kc— obtuvo menor MAE en cinco de los seis tramos de "
        "prueba. Por ello, Streamlit carga y muestra la estimación de cada modelo, "
        "pero solo permite que el modelo determine la dosis si su error de prueba "
        "fue menor que el físico: algodón BA08, fuera de ciclo en octubre. Los "
        "tres cultivos activos conservan el cálculo físico. Esta elección utiliza "
        "el tramo de prueba como control operativo; las cifras de la tabla ya no "
        "son una evaluación independiente del sistema completo seleccionado."
    )
    before(
        doc, chapter5,
        "Para mañana, primero se elige la estimación de ETc. Después, si AEMET "
        "publica una cantidad de lluvia horaria completa, se descuenta una única "
        "vez la fracción aprovechable hipotética indicada por el usuario; finalmente "
        "se aplica la eficiencia del sistema de riego. Una probabilidad de lluvia "
        "sin cantidad no se trata como cero. Si faltan uno o dos días intermedios "
        "de ETc, el modelo puede proyectarlos solo para comparación: esa extrapolación "
        "no está evaluada y nunca determina la dosis. No hay archivo histórico "
        "de previsiones AEMET alineadas ni riego real con que validar ahorro."
    )

    chapter52 = one_heading(doc, "5.2. Aplicación", "Heading 2")
    before(doc, chapter52, "5.1.7. Integración operativa de ETc y AEMET", "Heading 3")
    before(
        doc, chapter52,
        "La versión 1.1 del motor distingue la ETc estimada antes de lluvia, la "
        "precipitación AEMET disponible, la fracción efectiva supuesta, la necesidad "
        "neta y la dosis bruta. La predicción del modelo y el cálculo físico se "
        "comparan en la interfaz junto con sus MAE de prueba. El modelo solo "
        "interviene en el primer día si la observación SiAR del día anterior existe "
        "y su error histórico fue menor. Los días 2 a 7 mantienen la estimación "
        "física provisional. El CSV registra el método aplicado y, cuando procede, "
        "el modelo y la última fecha utilizada en su entrenamiento."
    )
    before(
        doc, chapter52,
        "El código de entrenamiento está en src/models/train_etc_for_dss.py, la "
        "inferencia en src/dss/predict_etc.py y el informe reproducible en "
        "docs/data_samples/etc_model_deployment_2026.json. Los seis artefactos "
        "pequeños de ETc se versionan para su carga en Streamlit. El cálculo "
        "físico permanece disponible si el modelo falta, envejece o no dispone "
        "de información reciente. La validación del capítulo 5.3 corresponde "
        "al motor anterior; no se presenta como evaluación prospectiva de esta "
        "integración reciente."
    )

    repair_index(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
