"""Sincroniza la memoria candidata con la recomendación de hoy del DSS."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/memoria/TFM_revision_final_contenido_2026-10-24.docx"
OUTPUT = ROOT / "docs/memoria/TFM_revision_final_2026-10-25.docx"


PARAGRAPHS = {
    192: (
        "El prototipo emite una recomendación para hoy o mañana y una planificación "
        "orientativa de hasta siete días. Calcula necesidad neta, dosis bruta según "
        "eficiencia y volumen cuando se indica la superficie. Para el día actual "
        "puede descontar una fracción supuesta de la lluvia AEMET prevista para "
        "las horas restantes, siempre que exista cantidad horaria utilizable; "
        "para mañana exige cobertura completa. Sin balance hídrico del suelo ni "
        "registro de riegos, la dosis de hoy representa la necesidad estimada "
        "del día completo y debe ajustarse si ya se ha regado. El sistema no "
        "determina el momento óptimo de riego ni acciona bombas o válvulas."
    ),
    281: (
        "Los módulos generan salidas reproducibles a partir de datos y parámetros "
        "conocidos. La suite reúne 103 pruebas definidas para extracción, protección "
        "de credenciales, escritura raw, normalización, calidad, variables, "
        "selección temporal, modelos de ETc y reglas del motor. Incluye controles "
        "para impedir el uso de información posterior a la consulta. Las figuras "
        "y resúmenes se regeneran desde los archivos del proyecto."
    ),
    515: (
        "Para el primer día, hoy o mañana, se estima primero la ETc previa a "
        "lluvia. Después, si AEMET ofrece milímetros horarios utilizables, se "
        "descuenta una sola vez la fracción aprovechable hipotética indicada en "
        "la interfaz; finalmente se aplica la eficiencia del sistema de riego. "
        "Hoy solo se suman horas posteriores a la consulta y mañana se exige "
        "cobertura de las 24 horas. Una probabilidad sin cantidad no equivale "
        "a lluvia cero. Si faltan uno o dos días intermedios de ETc, el modelo "
        "puede proyectarlos solo para comparación: esa extrapolación no determina "
        "la dosis. No hay previsiones AEMET históricas alineadas ni riego real "
        "con que validar ahorro."
    ),
    524: (
        "La Figura 5.1 muestra el flujo simplificado del motor retrospectivo "
        "de 2025: entradas, estimación diaria, reglas y salida. La lógica "
        "permanece en src/dss y puede usarse desde la línea de comandos o la "
        "interfaz. La ampliación operativa compara ETc del modelo con la vía "
        "física y aplica después la lluvia AEMET, como explica el apartado "
        "5.1.7; esos pasos posteriores no se dibujan en la figura."
    ),
    526: "Figura 5.1. Flujo simplificado del motor retrospectivo de 2025.",
    529: (
        "El motor distingue dos modos, que la interfaz denomina «Recomendación "
        "(hoy o mañana)» y «Revisión histórica (día cerrado)». La revisión "
        "histórica analiza fechas con variables observadas y aplica la "
        "expresión agronómica:"
    ),
    531: (
        "La revisión histórica permite comparar a posteriori una propuesta con "
        "la referencia SiAR; no acredita el riego aplicado en la parcela. En la "
        "recomendación para hoy o mañana todavía no se conoce la ET0 observada "
        "ni la precipitación efectiva del día objetivo. Para estimar su demanda "
        "se establecieron tres vías con un orden de preferencia fijo."
    ),
    534: (
        "En la validación retrospectiva de 2025 la lluvia prevista no se "
        "descontó por defecto: el motor avisaba y mantenía la dosis si no se "
        "había definido una fracción efectiva. La aplicación actual permite "
        "descontar una fracción supuesta cuando AEMET aporta milímetros "
        "horarios utilizables. La precipitación efectiva real depende de la "
        "intensidad, el suelo y las pérdidas por escorrentía o percolación "
        "(Allen et al., 1998); por ello, la cifra debe revisarse después de "
        "la lluvia. Si solo se conoce la probabilidad, no se descuenta agua."
    ),
    546: (
        "Durante el diseño original se valoró el modelo Ridge de necesidad "
        "neta del apartado 4.3 como motor principal. Se descartó porque no "
        "superó de forma estable a la persistencia y porque ante cambios de "
        "fase es menos transparente que la vía ET0 por Kc. Ese experimento "
        "es distinto del modelo Ridge de ETc para coliflor, que la aplicación "
        "actual muestra solo como contraste. También se descartó calcular "
        "ET0 futura con Hargreaves a partir de temperatura AEMET por falta "
        "de predicciones históricas para calibrarla (Allen et al., 1998). "
        "El motor admite una ET0 prevista explícita para una integración "
        "posterior. Un balance hídrico del suelo permitiría decidir cuándo "
        "regar, pero requiere capacidad de retención, profundidad radicular "
        "y riegos previos no disponibles en el piloto."
    ),
    556: (
        "La pantalla separa la selección del caso y la consulta de sus "
        "resultados. En el panel lateral se escogen comunidad, provincia, "
        "cultivo, estación y campaña entre CSV verificados. La consulta ofrece "
        "«Recomendación (hoy o mañana)» y «Revisión histórica (día cerrado)». "
        "Para un cultivo activo con observación SiAR reciente, el primer modo "
        "propone hoy por defecto y permite escoger mañana. El usuario ajusta "
        "horizonte, sistema de riego, eficiencia, superficie y dotación opcional. "
        "Al pulsar el botón se actualiza SiAR en memoria y se consulta AEMET "
        "para el caso vigente. Las tablas técnicas permanecen en desplegables."
    ),
    575: (
        "La aplicación es una demostración pública sin cuentas de usuario ni "
        "histórico de riego por parcela. Lee CSV verificados y puede actualizar "
        "en memoria el caso vigente desde SiAR. Para hoy suma solo las horas "
        "futuras de lluvia AEMET que tienen cantidad numérica utilizable; para "
        "mañana exige cobertura horaria completa. Una emisión antigua o sin "
        "milímetros útiles deja la dosis base sin descuento. La dosis de hoy "
        "representa el día completo: la app desconoce el riego y la lluvia ya "
        "ocurridos, por lo que el usuario debe ajustarla antes de aplicarla. "
        "No calcula ET0 futura a partir de la temperatura ni dispone de "
        "emisiones pasadas para evaluar AEMET. Los días posteriores al primero "
        "son orientativos y deben recalcularse. Tampoco se ha probado la "
        "interfaz con agricultores o técnicos (Rose et al., 2016)."
    ),
    652: (
        "La tercera limitación afecta a la predicción meteorológica. No existe "
        "una ET0 futura contrastada ni un archivo histórico de emisiones AEMET "
        "anteriores a cada fecha objetivo. Los resultados retrospectivos de "
        "2025 no miden la integración actual. La aplicación descuenta una "
        "fracción supuesta de la lluvia prevista cuando existen milímetros "
        "horarios utilizables: las horas restantes de hoy o las 24 horas de "
        "mañana. Esa fracción no está calibrada para la parcela y falta el "
        "balance de agua del suelo. Los días posteriores al primero usan una "
        "estimación provisional de ET0 y Kc. En la prueba original, el "
        "séptimo día perdió frente a la persistencia."
    ),
}


def replace_paragraph(paragraph, expected_start: str, value: str) -> None:
    if not paragraph.text.startswith(expected_start):
        raise ValueError(f"Ancla inesperada: {paragraph.text[:90]}")
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def main() -> None:
    document = Document(SOURCE)
    if len(document.paragraphs) != 724 or len(document.tables) != 30:
        raise ValueError("La estructura de la memoria candidata ha cambiado.")

    anchors = {
        192: "El prototipo emite",
        281: "Los módulos generan",
        515: "Para mañana",
        524: "La Figura 5.1 resume",
        526: "Figura 5.1.",
        529: "El motor distingue",
        531: "Este modo sirve",
        534: "La lluvia prevista",
        546: "Durante el diseño",
        556: "La pantalla separa",
        575: "La aplicación es",
        652: "La tercera limitación",
    }
    for index, value in PARAGRAPHS.items():
        replace_paragraph(document.paragraphs[index], anchors[index], value)

    index_figure = document.paragraphs[163]
    if not index_figure.text.startswith("Figura 5.1. Flujo de una consulta"):
        raise ValueError("No se encontró la entrada de la Figura 5.1 en el índice.")
    replace_paragraph(
        index_figure, "Figura 5.1.",
        "Figura 5.1. Flujo simplificado del motor retrospectivo de 2025.\t56",
    )

    rule = document.tables[20]
    replace_paragraph(
        rule.cell(5, 2).paragraphs[0], "Aviso de aplazamiento",
        "Aviso; en la consulta actual se descuenta la fracción supuesta "
        "si hay milímetros horarios utilizables",
    )
    scenario = document.tables[26]
    replace_paragraph(
        scenario.cell(2, 0).paragraphs[0], "Lluvia prevista de 12 mm",
        "Prueba retrospectiva: lluvia prevista de 12 mm con probabilidad del 80 %",
    )
    replace_paragraph(
        scenario.cell(2, 1).paragraphs[0], "Aviso sin descuento",
        "Aviso sin fracción efectiva configurada",
    )
    replace_paragraph(
        document.tables[27].cell(7, 1).paragraphs[0], "Apartado 5.3;",
        "Apartado 5.3; validación retrospectiva, 13 escenarios y "
        "103 pruebas automáticas definidas",
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
