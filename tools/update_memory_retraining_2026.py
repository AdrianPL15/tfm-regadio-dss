"""Incorpora el experimento de 2026 a una copia de la memoria de revisión."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/memoria/TFM_revision_coherencia_2026-10.docx"
OUTPUT = ROOT / "docs/memoria/TFM_revision_reentrenamiento_2026-10-23.docx"
FIGURE = ROOT / "docs/images/model_retraining_2026.png"


REPLACEMENTS = {
    "La evaluación empírica se limita al pimiento": (
        "El trabajo es una prueba de concepto de un sistema de apoyo a la decisión que conecta "
        "adquisición de datos, análisis, comparación de modelos y recomendación. El piloto "
        "original evaluó pimiento en AL01 La Mojonera durante 2025. La ampliación posterior "
        "reentrenó modelos independientes para seis combinaciones de cultivo y estación con "
        "las campañas disponibles hasta octubre de 2026. La interfaz pública permite escoger "
        "entre las combinaciones descargadas, pero no cualquier zona o cultivo de España. "
        "Estos ensayos no demuestran validez agronómica en parcela ni ahorro de agua."
    ),
    "La fuente de la evaluación es SiAR": (
        "La referencia de evaluación es la necesidad neta calculada por SiAR, no el riego "
        "realizado. Se usan nueve CSV oficiales de tres estaciones, seis cultivos y una o dos "
        "campañas según el caso. AEMET aporta una previsión municipal a la demostración actual "
        "cuando se consulta; sus emisiones no forman un histórico anterior a las fechas "
        "evaluadas y no se usan para entrenar. Los indicadores agregados del Boletín "
        "Hidrológico del MITECO no se interpretan como dotación de una parcela."
    ),
    "esa interfaz no se ha implementado en el piloto": (
        "La visión del producto contempla que el usuario seleccione libremente ubicación, "
        "cultivo y periodo para determinar las fuentes necesarias. La demo implementa una "
        "selección limitada a las combinaciones SiAR ya descargadas, no una extracción "
        "nacional arbitraria. El caso original de evaluación del motor sigue siendo pimiento "
        "en AL01 La Mojonera durante 2025. El CSV lo denomina Pimiento sin acreditar el "
        "sistema productivo; tratarlo como aire libre continúa siendo una hipótesis."
    ),
    "El ciclo completo permite validar la integración": (
        "El ciclo original de pimiento permite validar la integración y el motor en un caso "
        "acotado. El apartado 4.3 añade después una evaluación predictiva por separado en "
        "seis combinaciones y hasta dos campañas; el apartado 5.3 conserva su validación "
        "específica del motor en el piloto inicial. Esta ampliación no acredita un modelo "
        "transferible a otras estaciones ni una dosis correcta en parcela. La coherencia de "
        "unidades quedó confirmada en 5.3.1, con una diferencia máxima de 0,008 mm frente a "
        "la necesidad publicada por SiAR."
    ),
    "la muestra actual —un único ciclo, cultivo y estación—": (
        "En el experimento original de pimiento se mantiene la persistencia como "
        "referencia predictiva: ningún algoritmo de aprendizaje redujo de forma estable "
        "su MAE de validación. La ampliación de 2026 selecciona también persistencia en "
        "cinco de seis combinaciones; Ridge se selecciona en coliflor, aunque su error "
        "de prueba es ligeramente mayor que el de persistencia. Por tanto, los nuevos "
        "datos amplían los casos evaluados sin justificar todavía el uso de un modelo "
        "de aprendizaje común en el motor operativo."
    ),
    "ninguno de los cinco algoritmos de aprendizaje automático evaluados en el apartado 4.3": (
        "El capítulo anterior estableció dos referentes. La necesidad neta publicada "
        "por SiAR permite construir un cálculo agronómico explicable para jornadas "
        "cerradas. En el piloto original de pimiento, los cinco algoritmos de aprendizaje "
        "no mejoraron de forma estable la persistencia. El reentrenamiento posterior "
        "muestra resultados diferentes por cultivo, pero no valida un modelo compartido "
        "ni cambia el motor de recomendación. Este se apoya en ET0 reciente, Kc del "
        "calendario y reglas visibles para convertir necesidad neta en dosis."
    ),
    "En el primero se muestran el cultivo y la estación": (
        "La pantalla separa la selección y la consulta de los resultados. El panel lateral "
        "permite escoger comunidad, provincia, cultivo, estación y campaña entre los CSV "
        "verificados; también recoge fecha, horizonte, sistema de riego, eficiencia, "
        "superficie y dotación opcional. El área principal destaca primero la dosis diaria "
        "en L/m² y su volumen para la parcela, después el panorama orientativo y el gráfico. "
        "Las tablas de AEMET, el detalle diario y las reglas quedan en secciones desplegables. "
        "El botón de actualización consulta SiAR y, para un caso vigente, la previsión AEMET."
    ),
    "la suite completa contiene 61 pruebas": (
        "Los trece escenarios de validación del motor se complementan con pruebas automáticas "
        "sobre extracción, calidad, selección temporal y recomendaciones. La ampliación de "
        "2026 añade controles para las seis agrupaciones, las jornadas incompletas y los "
        "retardos entre campañas. La prueba del motor original sigue siendo distinta de la "
        "evaluación predictiva añadida; ambas utilizan la referencia calculada por SiAR."
    ),
    "el tercero ejecuta las sesenta y una pruebas": (
        "El primer comando produce el informe del DSS y sus figuras; el segundo genera la "
        "Figura 5.1 y el tercero ejecuta la suite de pruebas. El reentrenamiento adicional "
        "se reproduce con python -m src.models.retrain_cases_2026. La captura de la "
        "aplicación se obtiene desde el enlace público de Streamlit o arrancando "
        "streamlit run src/dss/app_streamlit.py."
    ),
    "faltan datos hidrológicos y de parcela, diversidad de campañas y cultivos": (
        "El balance de objetivos mantiene las reservas indicadas en la Tabla 6.1. La nueva "
        "evaluación incorpora seis combinaciones y hasta dos campañas, pero sigue faltando "
        "diversidad de fincas y años para demostrar transferencia. Tampoco se dispone de "
        "riego aplicado, suelo o respuesta productiva, ni de un balance hídrico para decidir "
        "el momento de riego. El motor admite eficiencia y dotación, aunque no puede "
        "interpretar una recomendación anterior como riego efectivamente realizado."
    ),
    "Las sesenta y una pruebas automáticas se han ejecutado": (
        "La cuarta aportación es la reproducibilidad. El repositorio conserva los CSV SiAR "
        "acotados, el protocolo de seis enfoques, las predicciones y las métricas por caso, "
        "además del motor y la demo pública. El reentrenamiento se ejecuta desde un módulo "
        "único y genera artefactos locales excluidos de Git. Las pruebas automáticas del "
        "experimento ampliado verifican calidad, fechas, retardos y ausencia de variables "
        "meteorológicas del propio día; el README documenta cómo repetirlo."
    ),
    "Todo se ha evaluado con una estación, un cultivo y una campaña": (
        "La segunda limitación es de validez externa. La evaluación del DSS del capítulo 5 "
        "sigue centrada en pimiento AL01 en 2025. La evaluación predictiva ampliada reúne "
        "seis combinaciones de tres estaciones y una o dos campañas, pero cada prueba final "
        "pertenece a un tramo temporal concreto de 2026. Coliflor y alcachofa presentan "
        "mayor error en prueba que en validación. No se ha demostrado transferencia a otras "
        "fincas, sistemas de producción ni años posteriores."
    ),
    "La aplicación, por su parte, trabaja con un caso fijo": (
        "La cuarta limitación es funcional. El sistema propone una dosis orientativa, pero "
        "no decide cuándo regar porque no mantiene un balance hídrico del suelo. No sabe si "
        "una recomendación anterior se aplicó y no la descuenta después. La reducción por "
        "dotación tampoco distribuye el déficit según la fase del cultivo. La app permite "
        "elegir solo entre los casos cargados, no conserva un histórico de consultas y no "
        "se ha probado con agricultores ni técnicos."
    ),
    "Ninguno de los algoritmos de aprendizaje superó de forma estable a esta referencia sencilla": (
        "La segunda aportación es una comparación temporal reproducible. En el piloto "
        "original se midieron persistencia, Ridge, Random Forest, KNN, SVR y XGBoost "
        "sobre la referencia calculada por SiAR, sin ET0, ETc ni precipitación efectiva "
        "del mismo día entre los predictores. El modelo se eligió antes de abrir la prueba "
        "final. La ampliación posterior reentrenó por separado seis combinaciones: "
        "persistencia fue elegida en cinco y Ridge en coliflor, donde su MAE final "
        "(0,892 mm/día) no superó al de persistencia (0,885). Estas pruebas no indican "
        "que un algoritmo complejo mejore de forma general la estimación. La evaluación "
        "del método físico de capítulo 5 reutiliza fechas del piloto y no constituye "
        "otra prueba independiente."
    ),
    "La demostración local sin nuevas consultas": (
        "La demostración pública se encuentra en "
        "https://tfm-regadio-dss-adrianpl15.streamlit.app/ y también puede iniciarse con "
        "streamlit run src/dss/app_streamlit.py. Permite seleccionar casos CSV verificados "
        "y consultar SiAR y AEMET para un cultivo vigente. La alternativa local sin nuevas "
        "consultas sigue descrita en docs/DEMO_TFM.md. Las recomendaciones son orientativas: "
        "la evaluación usa la necesidad calculada por SiAR, no riego ni rendimiento de campo."
    ),
}


def replace_paragraphs(doc: Document) -> None:
    for fragment, replacement in REPLACEMENTS.items():
        matches = [p for p in doc.paragraphs if fragment in p.text]
        if len(matches) != 1:
            raise ValueError(f"Se esperaban 1 coincidencia para {fragment!r}: {len(matches)}")
        matches[0].text = replacement


def insert_before(anchor, paragraph):
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_before(doc: Document, anchor, text: str, style: str = "Normal"):
    paragraph = doc.add_paragraph(text, style=style)
    return insert_before(anchor, paragraph)


def make_table(doc: Document, anchor) -> None:
    rows = [
        ["Caso y campañas", "Días válidos", "Prueba", "Modelo", "MAE validación / prueba"],
        ["Pimiento AL01 · 2025–26", "300", "45", "Persistencia", "0,501 / 0,390"],
        ["Coliflor AL01 · 2026", "144", "22", "Ridge", "0,264 / 0,892"],
        ["Maíz Z06 · 2025–26", "366", "55", "Persistencia", "0,782 / 0,814"],
        ["Algodón BA08 · 2025–26", "365", "55", "Persistencia", "0,626 / 0,457"],
        ["Alfalfa Z06 · 2026", "236", "36", "Persistencia", "0,697 / 0,786"],
        ["Alcachofa BA08 · 2026", "173", "26", "Persistencia", "0,475 / 0,896"],
    ]
    table = doc.add_table(rows=len(rows), cols=5, style="TableNormal")
    table.autofit = False
    widths = [Inches(1.62), Inches(0.68), Inches(0.59), Inches(1.07), Inches(1.98)]
    for row_idx, values in enumerate(rows):
        for col_idx, value in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            cell.width = widths[col_idx]
            cell.vertical_alignment = 1
            for p in cell.paragraphs:
                if col_idx > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8)
                    if row_idx == 0:
                        run.bold = True
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = tc_pr.find(qn("w:tcBorders"))
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tc_pr.append(borders)
            for edge in ("top", "left", "bottom", "right"):
                node = OxmlElement(f"w:{edge}")
                node.set(qn("w:val"), "single")
                node.set(qn("w:color"), "D9D9D9")
                node.set(qn("w:sz"), "4")
                borders.append(node)
            if row_idx == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF5")
                tc_pr.append(shading)
    anchor._p.addprevious(table._tbl)


def insert_extension(doc: Document) -> None:
    anchors = [p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.startswith("5. Sistema")]
    if len(anchors) != 1:
        raise ValueError("No se encontró el inicio del capítulo 5.")
    anchor = anchors[0]
    add_before(doc, anchor, "4.3.10. Reentrenamiento con datos de 2026", "Heading 3")
    add_before(
        doc, anchor,
        "El experimento original se amplió a los nueve CSV oficiales disponibles, agrupados "
        "en seis combinaciones de cultivo y estación. Pimiento, maíz y algodón incluyen "
        "2025 y 2026; coliflor, alfalfa y alcachofa solo tienen la campaña parcial de 2026. "
        "De 1776 filas originales, 1584 jornadas completas y con Kc positivo se emplearon "
        "en la evaluación. Los huecos no se sustituyeron por necesidad cero y los días "
        "fuera del ciclo no contaron en las métricas."
    )
    add_before(
        doc, anchor,
        "Se mantuvieron los nueve predictores y los seis enfoques anteriores. Los retardos "
        "se generaron por fecha de calendario antes de filtrar jornadas: el primer día de "
        "2026 no hereda la última necesidad de 2025. En cada caso se reservó el 15 % "
        "cronológicamente final de los días activos como prueba; la elección se hizo por "
        "el MAE medio de tres cortes anteriores. Después de evaluar, se reajustó el enfoque "
        "elegido con todos los datos válidos y se guardó el artefacto local, sin volver a "
        "evaluarlo sobre las mismas filas."
    )
    add_before(
        doc, anchor,
        "Tabla 4.15. Reentrenamiento por cultivo y estación con datos hasta el 22 de octubre de 2026. "
        "Errores en mm/día.",
        "Índice tabla TFM",
    )
    make_table(doc, anchor)
    add_before(
        doc, anchor,
        "La persistencia fue elegida en cinco casos. Ridge obtuvo el menor error de "
        "validación en coliflor, pero en su prueba final registró 0,892 mm/día, frente a "
        "0,885 de persistencia. Coliflor y alcachofa muestran además un aumento notable "
        "del error entre validación y prueba. Los periodos y tamaños de prueba difieren "
        "de los del experimento de pimiento de 2025; sus cifras no representan una "
        "mejora comparable ni una validación agronómica en campo."
    )
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(FIGURE), width=Inches(5.15))
    insert_before(anchor, picture)
    add_before(
        doc, anchor,
        "Figura 4.8. Error del modelo elegido en validación temporal y prueba final para seis casos. "
        "Fuente: elaboración propia a partir de los CSV SiAR.",
        "Índice figura TFM",
    )
    add_before(
        doc, anchor,
        "El reentrenamiento amplía la evidencia, pero la etiqueta es una necesidad "
        "calculada por SiAR, no riego aplicado. La persistencia predice a un día con "
        "el dato anterior conocido. Sin previsiones AEMET archivadas antes de cada fecha "
        "objetivo no se puede medir su aportación. Streamlit sigue usando el motor de "
        "reglas, no estos seis modelos."
    )


def insert_index_entries(doc: Document) -> None:
    entries = [
        ("4.3.9. Reproducibilidad y evidencias", "4.3.10. Reentrenamiento con datos de 2026\t52"),
        ("Tabla 4.14. Resultados sobre el periodo de prueba reservado.",
         "Tabla 4.15. Reentrenamiento por cultivo y estación con datos hasta el 22 de octubre de 2026.\t53"),
        ("Figura 4.7. Predicciones y errores diarios de los seis enfoques en el tramo de prueba.",
         "Figura 4.8. Error del modelo elegido en validación temporal y prueba final para seis casos.\t53"),
    ]
    for old_text, new_text in entries:
        matches = [p for p in doc.paragraphs if p.text.startswith(old_text) and "\t" in p.text]
        if len(matches) != 1:
            raise ValueError(f"No se encontró una única entrada de índice: {old_text}")
        anchor = matches[0]
        new = doc.add_paragraph(new_text, style=anchor.style)
        for stop in anchor.paragraph_format.tab_stops:
            new.paragraph_format.tab_stops.add_tab_stop(
                stop.position, stop.alignment, stop.leader
            )
        anchor._p.addnext(new._p)
def main() -> None:
    doc = Document(SOURCE)
    replace_paragraphs(doc)
    insert_extension(doc)
    insert_index_entries(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
