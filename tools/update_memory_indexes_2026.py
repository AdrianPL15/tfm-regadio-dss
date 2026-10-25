"""Synchronize the manually typeset indices with the rendered final memory."""

import argparse
from pathlib import Path
import re
import unicodedata

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCX = ROOT / "docs/memoria/TFM_revision_final_contenido_2026-10-24.docx"
DEFAULT_PDF = ROOT / "docs/memoria/TFM_revision_final_2026-10-24.pdf"

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
parser.add_argument("--pdf", type=Path, default=DEFAULT_PDF)
args = parser.parse_args()
DOCX = args.docx
PDF = args.pdf


def normalized(value):
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"\W+", "", value).lower()


document = Document(DOCX)
pages = [normalized(page.extract_text() or "") for page in PdfReader(PDF).pages]
changed = []

for index in [*range(22, 119), *range(122, 152), *range(154, 168)]:
    paragraph = document.paragraphs[index]
    title = paragraph.text.rsplit("\t", 1)[0]
    if index == 155:
        title = next(
            p.text for p in document.paragraphs[168:] if p.text.startswith("Figura 4.1.")
        )
        paragraph.runs[0].text = title
    stem = normalized(title)[: min(65, len(normalized(title)))]
    matches = [page + 1 for page, value in enumerate(pages) if page >= 8 and stem in value]
    if len(matches) != 1:
        raise RuntimeError(f"Index {index}: ambiguous page match {matches}: {title}")
    nodes = list(paragraph._p.iter(qn("w:t")))
    if not nodes or not re.fullmatch(r"\d+", nodes[-1].text or ""):
        raise RuntimeError(f"Index {index}: missing numeric final node: {paragraph.text}")
    previous = nodes[-1].text
    nodes[-1].text = str(matches[0])
    if previous != nodes[-1].text:
        changed.append((index, previous, nodes[-1].text))

document.save(DOCX)
print(f"Updated {len(changed)} page references in {DOCX}")
